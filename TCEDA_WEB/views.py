from django.shortcuts import render
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.writer.excel import save_virtual_workbook  # Create your views here.
import requests
from openpyxl.chart import BarChart
import collections
from openpyxl.chart import (
    ScatterChart,
    Reference,
    Series,
)

merced = '05000US06047'
tuolumne = '05000US06109'
us = '01000US'
california = '04000US06'
location_dict = {
    merced: 'Merced',
    tuolumne: 'Tuolumne',
    us: 'United States',
    california: 'California',
}


def population():
    headers = {
        'X-App-Token': 'cQovpGcdUT1CSzgYk0KPYdAI0'
    }
    url = 'https://api.opendatanetwork.com/data/v1/values?app_token=cQovpGcdUT1CSzgYk0KPYdAI0&forecast=8&variable=demographics.population.count&entity_id=0500000US06109&format=null&describe=true'
    r = requests.get(url, headers=headers)
    x = True
    out = collections.OrderedDict()
    header = {
        'year': 'Year',
        'projected': 'Projected',
        'population': 'Population',
    }
    for year in r.json()['data']:
        if x:
            x = False
            continue
        else:
            date = year[0]
            population = year[1]
            projected = year[2]
            year_data = "Year %s: %s" % (date, int(population))
            if projected:
                year_data += " (Projected)"

            out.update(
                {
                    date: {
                        'date': date,
                        'population': population,
                        'projected': projected,
                    }
                }
            )
    out_dict = {
        'header': header,
        'data': out,
    }
    return out_dict


def household_income():
    url = 'https://api.datausa.io/api/?sort=desc&show=geo&required=income%2Cincome_moe&sumlevel=all&year=all&geo=01000US%2C04000US06%2C05000US06109%2C05000US06047'
    r = requests.get(url)
    header = {}
    header.update({
        'year': 'Year',
        'location': 'Location',
        'income': 'Median Household Income',
    })
    out = {}
    for location in r.json()['data']:
        year = location[0]
        location_code = location[1]
        income = location[2]
        if year not in out:
            out.update(
                {
                    year: {
                        location_code: {
                            'location': location_dict[location_code],
                            'income': income,
                            'year': year,
                        }
                    }
                }
            )
        else:
            out[year].update(
                {
                    location_code: {
                        'location': location_dict[location_code],
                        'income': income,
                        'year': year,
                    }
                }
            )
    out_dict = {
        'header': header,
        'data': out,
    }
    return out_dict


def get_householdIncome():
    workbook = Workbook()
    worksheet = workbook.active
    household_data = household_income()
    worksheet.append(list(household_data['header'].values()))
    for year, data in household_data['data'].iteritems():
        for x, y in data.iteritems():
            worksheet.append(y.values())

    chart1 = BarChart()
    chart1.type = "col"
    chart1.style = 10
    chart1.title = "Bar Chart"
    chart1.y_axis.title = household_data['header']['income']
    chart1.x_axis.title = household_data['header']['location']

    data = Reference(worksheet, min_col=3, min_row=1, max_row=((len(location_dict) + 1) * 2) - 1,
                     max_col=3)
    cats = Reference(worksheet, min_col=1, min_row=2, max_row=(len(location_dict) + 1) * 2,
                     max_col=2)
    chart1.add_data(data, titles_from_data=True)
    chart1.set_categories(cats)
    chart1.shape = 4
    worksheet.add_chart(chart1, "A15")

    response = HttpResponse(content=save_virtual_workbook(workbook),
                            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=household.xlsx'
    return response


def get_population():
    workbook = Workbook()
    worksheet = workbook.active
    population_data = population()
    worksheet.append(list(population_data['header'].values()))
    for date, data in population_data['data'].iteritems():
        worksheet.append(data.values())

    chart = ScatterChart()
    chart.scatterStyle = 'marker'
    chart.title = "Tuolumne County Population"
    chart.style = 13
    chart.x_axis.title = population_data['header']['year']
    chart.y_axis.title = population_data['header']['population']

    xvalues = Reference(worksheet, min_col=3, min_row=1, max_row=len(population_data['data']) + 1)
    values = Reference(worksheet, min_col=1, min_row=2, max_row=len(population_data['data']) + 1)
    series = Series(xvalues, values, title_from_data=True)
    chart.series.append(series)

    worksheet.add_chart(chart, "A10")
    response = HttpResponse(content=save_virtual_workbook(workbook),
                            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=population.xlsx'

    return response


def index(request):
    if request.POST:
        if request.POST.get('name') == 'getPopulation':
            return get_population()
        if request.POST.get('name') == 'gethouseholdIncome':
            return get_householdIncome()
    return render(request, 'TCEDA_WEB/index.html', {})
